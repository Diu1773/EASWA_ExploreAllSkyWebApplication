# -*- coding: utf-8 -*-
"""v20 — 강조·색·화면 (2026-09-11).

사장님 지시 —
  *「easwa링크를 웹기능파트에 넣고」*
  *「easwa구조는 걍 도식없애고, 홈페이지 화면이랑 모듈안에 하나 넣고 보자」*
  *「여기 이미 두개사진 빼고 그냥 연구배경2 도식에서 중간 절차들에 X도식 넣고
    우리 탐구활동들이 이어지도록 만드는 도식을 넣자」*
  *「중요한 단어 글에는 굵게 강조한다던가」*

  ① `put()` 이 «…» 를 굵은 강조 런으로 그린다. 이미 원고 곳곳에 쓰던 표시라
     새 문법을 들이지 않고 강조가 붙는다. 본문색 문단에서는 파랑으로도 바꾼다
  ② 7장 — 화면 둘을 빼고 `fig_bridge.png` 를 넣는다
  ③ 12장 — 논문 도식을 빼고 «홈 화면 + 모듈 안 화면» 둘에 EASWA 주소를 붙인다
  ④ 중요한 낱말에 «…» 를 더 단다
"""
import io
import os

HERE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(HERE, "build.py")
s = io.open(P, encoding="utf-8").read()


def swap(old, new, label):
    global s
    assert s.count(old) == 1, "%s — %d 곳" % (label, s.count(old))
    s = s.replace(old, new)


# ══ ① «…» 를 굵은 강조로 ══════════════════════════════════════════════
swap("""import os
from pptx import Presentation""",
     """import os
import re

from pptx import Presentation""", "import re")

swap("""RULEC = C(0xD9, 0xD9, 0xD9)
WHITE = C(0xFF, 0xFF, 0xFF)""",
     """RULEC = C(0xD9, 0xD9, 0xD9)
WHITE = C(0xFF, 0xFF, 0xFF)
ORANGE = C(0xE8, 0x72, 0x2A)   # EASWA 앱 화면에서 뽑은 값 (2026-09-11 실측)
EMPH = C(0x2E, 0x75, 0xB6)     # «…» 로 묶은 낱말 — 본문 문단에서만 색을 바꾼다""",
     "팔레트")

swap('''    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, F
    r.font.color.rgb = color if color is not None else BODY
    return p''',
     '''    base = color if color is not None else BODY
    # «…» 로 묶은 데는 굵게 준다. 원고에서 이미 쓰던 표시라 새 문법이 아니다.
    for part in re.split("(«[^»]*»)", text):
        if not part:
            continue
        hot = part.startswith("«") and part.endswith("»")
        r = p.add_run()
        r.text = part
        r.font.size, r.font.name = Pt(size), F
        r.font.bold = True if hot else bold
        r.font.color.rgb = EMPH if (hot and base in (BODY, GREY)) else base
    return p''', "put 강조")

# 강조 상자 — 굵은 테두리만 두고 안은 비운다
swap('''def bullets(sl, x, y, w, items, size=19, gap=13):''',
     '''def chip(sl, x, y, w, h, text, size=16, color=None, fill=None):
    """강조 상자 하나. 테두리 색으로 눈을 끈다."""
    color = color if color is not None else ORANGE
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.adjustments[0] = 0.18
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = color
    sh.line.width = Pt(1.8)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    put(tf, text, size, color, True, first=True, align=PP_ALIGN.CENTER)
    return sh


def bullets(sl, x, y, w, items, size=19, gap=13):''', "chip")

# ══ ② 7장 — 화면 둘을 빼고 도식을 넣는다 ══════════════════════════════
swap('''pic(sl, "easwa_home.png", M, y + 0.04, W - 2 * M, 3.05, root=HERE, top=True)
bullets(sl, M, y + 3.26, W - 2 * M, [''',
     '''pic(sl, "fig_bridge.png", M, y + 0.02, W - 2 * M, 3.22, root=HERE, top=True)
bullets(sl, M, y + 3.30, W - 2 * M, [''', "7장 그림")

swap('''    ("① 공개 아카이브 자료를 그대로 쓰되, 검색·내려받기·형식 변환·반복 계산은 플랫폼이 맡음",
     BODY, False),
    ("② 자료 구조가 다른 여러 주제를 «같은» 탐구 흐름 안에 둠 — 주제마다 다른 도구로 "
     "흩어지지 않게", BODY, False),
    ("③ 분석을 밖으로 넘기지 않음. 대신 분석 조건을 학습자가 정하고 그 조건을 화면에 남김",
     ACC, True),
], size=16, gap=9)''',
     '''    ("① 공개 아카이브 자료를 그대로 쓰되, 검색·내려받기·형식 변환·반복 계산은 "
     "«플랫폼이 맡음»", BODY, False),
    ("② 자료 구조가 다른 여러 주제를 «같은 탐구 흐름» 안에 둠 — 주제마다 다른 도구로 "
     "흩어지지 않게", BODY, False),
    ("③ 분석을 밖으로 넘기지 않음. 대신 «분석 조건을 학습자가 정하고» 그 조건을 화면에 남김",
     ACC, True),
], size=15, gap=7)''', "7장 글줄")

# ══ ③ 12장 — 도식을 빼고 화면 둘 + 주소 ═══════════════════════════════
i = s.index('y = title(sl, "EASWA의 구조",')
i = s.rindex("# ═════", 0, i)
j = s.index("# ═════", i + 10)
struct = '''# ═════ EASWA 의 구조 — 화면 둘 ═══════════════════════════════════════
sl = S()
y = title(sl, "EASWA의 구조",
          "자료가 다른 세 모듈에 같은 일곱 단계를 적용함")
PH = 3.34
_a = Image.open(os.path.join(FIG, "home_hero.png")).size
_b = Image.open(os.path.join(FIG, "step4_analysis.png")).size
_wa, _wb = PH * _a[0] / _a[1], PH * _b[0] / _b[1]
_x = (W - (_wa + _wb + 0.34)) / 2
pic(sl, "home_hero.png", _x, y + 0.06, _wa, PH, top=True)
pic(sl, "step4_analysis.png", _x + _wa + 0.34, y + 0.06, _wb, PH, top=True)
f = tb(sl, _x, y + 3.46, _wa, 0.3, PP_ALIGN.CENTER)
put(f, "홈 화면 — 세 모듈을 «대표 탐구 질문»과 함께 제시", 14, BODY, first=True,
    align=PP_ALIGN.CENTER)
g = tb(sl, _x + _wa + 0.34, y + 3.46, _wb, 0.3, PP_ALIGN.CENTER)
put(g, "모듈 안 — 분석 실행과 «조건·품질 지표»를 같은 화면에", 14, BODY, first=True,
    align=PP_ALIGN.CENTER)
chip(sl, 3.30, H - 1.28, 6.73, 0.52, "직접 열어 보실 수 있음 — easwa-webapp.onrender.com", 17)
f = tb(sl, M, H - 0.62, W - 2 * M, 0.5)
put(f, "설치·로그인 없이 브라우저에서 바로 열림. 다만 세부 구현과 검토 범위는 같지 않음 — "
       "사용자 검토를 받은 것은 «식현상 모듈뿐»임", 14, WARN, True, first=True)

'''
s = s[:i] + struct + s[j:]

# ══ ④ 중요한 낱말에 «…» 를 더 단다 ═══════════════════════════════════
MARK = [
    ("어렵다는 점을 생각하면 오히려 학교 천문탐구의 자료 기반임",
     "어렵다는 점을 생각하면 오히려 «학교 천문탐구의 자료 기반»임"),
    ("그래서 준비 절차는 EASWA 가 맡고, 학습자는 자료 확인과 조건 선택과 해석에 ",
     "그래서 «준비 절차»는 EASWA 가 맡고, 학습자는 «자료 확인과 조건 선택과 해석»에 "),
    ("→ 탐구 질문에서 출발하는 진입 경로가 없고, 분석용 원자료는 따로 구해야 함",
     "→ «탐구 질문에서 출발하는 진입 경로»가 없고, 분석용 원자료는 따로 구해야 함"),
    ("→ 탐구활동이 분류에 머묾. 대상 선택도 측광도 모델 적합도 없음",
     "→ 탐구활동이 «분류에 머묾». 대상 선택도 측광도 모델 적합도 없음"),
    ("자료를 고르는 일 자체가 탐구의 일부임",
     "«자료를 고르는 일» 자체가 탐구의 일부임"),
    ("조건이 값을 바꾼다는 것 자체가 학습 내용임",
     "«조건이 값을 바꾼다»는 것 자체가 학습 내용임"),
    ("과제는 값 맞히기가 아니라 차이 설명임",
     "과제는 값 맞히기가 아니라 «차이 설명»임"),
    ("→ 실행 부담은 낮아졌고 해석 지원이 남음",
     "→ 실행 부담은 낮아졌고 «해석 지원»이 남음"),
    ("→ 자동화만으로는 용어와 해석이 해결되지 않음",
     "→ 자동화만으로는 «용어와 해석»이 해결되지 않음"),
    ("→ ①②가 먼저임. 해석 지원이 남은 과제로 확인된 이상 학생이 실제로 어디서 ",
     "→ «①②가 먼저»임. 해석 지원이 남은 과제로 확인된 이상 학생이 실제로 어디서 "),
    ("실행을 돕는 일과 해석을 돕는 일은 다른 과제임",
     "«실행을 돕는 일»과 «해석을 돕는 일»은 다른 과제임"),
]
for old, new in MARK:
    swap(old, new, "강조 — " + old[:18])

io.open(P, "w", encoding="utf-8", newline="\n").write(s)
print("강조 문법 · 7장 도식 · 12장 화면 둘 · 낱말 강조 %d곳" % len(MARK))
