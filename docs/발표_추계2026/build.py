# -*- coding: utf-8 -*-
"""한국지구과학회 추계 AS2(천문) 구두발표 — 15분 슬롯(발표 12 · 질의 3).

v2 (2026-09-10 밤) — 사장님 지시로 구성을 바꿨다.
  늘림 : 왜 만들었나(동기) · 웹 기능 셋 · 바이브 코딩과 그래서 한 검증 · 검토 결과
  줄임 : 사례분석 절차와 설계 원리 도출 같은 «방법» 서술

  python deck/build.py

흰 바탕이다. 2026-09-10 밤에 EASWA 다크 서식으로도 만들어 봤으나 사장님이
「눈에 안 들어온다」고 물리셨다. 강의실 프로젝터에서는 흰 바탕이 맞다.
"""
import os
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

FIG = r"C:/Users/bmffr/Desktop/Me/ERP2026_Cosmos/원고_그림"
HERE = os.path.dirname(os.path.abspath(__file__))
DEST = r"C:/Users/bmffr/Desktop/Me/ERP2026_Cosmos/추계학술발표회"

OUT = os.path.join(DEST, "EASWA_지구과학회_구두발표_2026-09-18.pptx")

BG = C(0xFF, 0xFF, 0xFF)
ACC = C(0x1F, 0x4E, 0x79)
ACC2 = C(0x2E, 0x75, 0xB6)
GREY = C(0x70, 0x70, 0x70)
WARN = C(0xC0, 0x00, 0x00)
BODY = C(0x1A, 0x1A, 0x1A)
HEAD = C(0x1F, 0x4E, 0x79)
TBLHEAD = C(0x1F, 0x4E, 0x79)
TBLTEXT = C(0x1A, 0x1A, 0x1A)
RULEC = C(0xD9, 0xD9, 0xD9)
WHITE = C(0xFF, 0xFF, 0xFF)

F = "맑은 고딕"
W, H = 13.333, 7.5
M = 0.62

prs = Presentation()
prs.slide_width, prs.slide_height = I(W), I(H)
BLANK = prs.slide_layouts[6]


def S():
    return prs.slides.add_slide(BLANK)


def tb(sl, x, y, w, h, align=PP_ALIGN.LEFT):
    s = sl.shapes.add_textbox(I(x), I(y), I(w), I(h))
    f = s.text_frame
    f.word_wrap = True
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    f.paragraphs[0].alignment = align
    return f


def put(f, text, size, color=None, bold=False, space_after=0, first=False,
        align=None, line=1.25):
    p = f.paragraphs[0] if first else f.add_paragraph()
    if align is not None:
        p.alignment = align
    p.line_spacing = line
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, F
    r.font.color.rgb = color if color is not None else BODY
    return p


def title(sl, text, sub=None):
    """행동 제목 — 요지를 문장으로 적는다."""
    f = tb(sl, M, 0.42, W - 2 * M, 1.0)
    put(f, text, 25, HEAD, True, first=True, line=1.15)
    if sub:
        g = tb(sl, M, 1.42, W - 2 * M, 0.4)
        put(g, sub, 16, C(0x44, 0x44, 0x44), first=True)
        return 2.00
    return 1.68


def cite(sl, text):
    f = tb(sl, M, H - 0.54, W - 2 * M, 0.34)
    put(f, text, 12, GREY, first=True)


def pic(sl, name, x, y, maxw, maxh, root=FIG, top=False):
    p = os.path.join(root, name)
    iw, ih = Image.open(p).size
    sc = min(maxw / iw, maxh / ih)
    w, h = iw * sc, ih * sc
    dy = 0 if top else (maxh - h) / 2
    sl.shapes.add_picture(p, I(x + (maxw - w) / 2), I(y + dy), I(w), I(h))


# 화면 캡처를 통째로 넣으면 강의실 뒤에서 글자가 안 읽힌다(2026-09-10 확인).
# 내비게이션과 제목 영역을 잘라내고 내용만 남긴다. 원본은 2880×2300 이다.
CROP = {
    "step0_entry.png":     (225, 730, 2700, 1700),   # 탐구 질문 + 시뮬레이션 + 광도곡선
    "step3_conditions.png": (255, 800, 2700, 2300),  # Step 3 머리부터 구경·배경 카드까지
}


def crop(name):
    """잘라 둔 사본을 만들고 그 이름을 준다. 이미 있으면 그대로 쓴다."""
    box = CROP.get(name)
    if not box:
        return name
    out = "_crop_" + name
    dst = os.path.join(HERE, out)
    if not os.path.exists(dst):
        Image.open(os.path.join(FIG, name)).crop(box).save(dst)
    return out


def bullets(sl, x, y, w, items, size=19, gap=13):
    f = tb(sl, x, y, w, H - y - 0.85)
    for i, (t, c, b) in enumerate(items):
        put(f, t, size, c, b, space_after=gap, first=(i == 0))


def rule(sl, y, x=M, w=None, color=None):
    w = w if w else W - 2 * M
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Emu(12700))
    s.fill.solid()
    s.fill.fore_color.rgb = color if color is not None else RULEC
    s.line.fill.background()
    s.shadow.inherit = False


# ═════ 1. 표제 ════════════════════════════════════════════════════════
sl = S()
f = tb(sl, M, 2.20, W - 2 * M, 1.7)
put(f, "천문 탐구 웹 플랫폼 EASWA의", 36, HEAD, True, first=True, line=1.2)
put(f, "개발과 현장 적용 방안", 36, HEAD, True, line=1.2)
rule(sl, 4.18, M, 3.2, ACC)
f = tb(sl, M, 4.48, W - 2 * M, 1.5)
put(f, "박민준 · 손정주", 20, BODY, True, first=True, space_after=6)
put(f, "한국교원대학교 지구과학교육과", 16, GREY, space_after=18)
put(f, "2026년 한국지구과학회 추계학술발표회 · AS2 천문 · 9월 18일 11:15", 14, GREY)

# ═════ 2. 왜 만들었나 ═════════════════════════════════════════════════
sl = S()
y = title(sl, "연구 배경",
          "자료는 공개되어 있지만 수업에서 쓰려면 준비 절차가 앞선다")
bullets(sl, M, y, W - 2 * M, [
    ("MAST · Gaia · KMTNet 등 공개 아카이브는 영상, 시계열 측광, 측성·측광 카탈로그를 상시 제공한다. "
     "자료는 이미 충분하다.", BODY, False),
    ("그런데 수업에서 쓰려면 검색 → 내려받기 → 형식 변환 → 코딩 → 반복 계산을 먼저 지나야 한다. "
     "이 절차는 학습 목표와 무관하다.", BODY, False),
    ("코딩을 쓰는 교사교육 프로그램에서도 현직·예비교사 모두 파이썬 코딩을 학교 적용의 "
     "가장 큰 어려움으로 들었다.", BODY, False),
    ("", BODY, False),
    ("그래서 만들었다. 준비는 도구가 지고, 자료 확인과 조건 선택과 해석은 학습자가 한다.",
     ACC, True),
], gap=15)
cite(sl, "교육부 (2022) 과학과 교육과정 · Wong et al. (2026)")

# ═════ 3. 기존 서비스 (방법은 줄이고 결과만) ══════════════════════════
sl = S()
y = title(sl, "기존 서비스 분석",
          "네 서비스 모두 탐구 흐름은 사용자 몫이었다")
pic(sl, "case_stage1_entry.png", M, y, 7.1, H - y - 0.9)
bullets(sl, M + 7.5, y + 0.2, W - M - 7.5 - M + 0.4, [
    ("SIMBAD · VizieR — 천체명·좌표·카탈로그 질의에서 시작한다.", BODY, False),
    ("WorldWide Telescope · ESASky — 시각 탐색과 통합 자료 접근이 중심이다.", BODY, False),
    ("네 곳 모두 단계별 질문과 기록 활동을 서비스 안에 두지 않았다.", ACC, True),
], size=17, gap=16)
cite(sl, "학교 활용 관점의 연구자 워크스루 (2026). 분석 기준과 결과는 논문 표 2·3.")

# ═════ 4. EASWA 개요 ══════════════════════════════════════════════════
sl = S()
y = title(sl, "EASWA의 구조",
          "자료가 다른 세 모듈에 같은 일곱 단계를 적용하였다")
f = tb(sl, M, y - 0.08, W - 2 * M, 0.34)
put(f, "탐구 주제 소개 → 대상 선택 → 자료 확인 → 분석 준비 → 분석·시각화 → "
       "기준값 비교 → 해석·기록", 13, GREY, first=True)
y += 0.30
rows = [
    ("모듈", "공공 자료", "분석 구조", "학습자가 정하는 것"),
    ("TESS 외계행성 식현상", "TESS FFI 컷아웃", "구경·차등측광 → 식현상 모델 적합", "측광 구경 · 배경 · 비교성"),
    ("KMTNet 미시중력렌즈", "관측소별 공개 측광표", "다지점 병합 → 점렌즈 모델 적합", "적합 대상 · 조건 확인"),
    ("Gaia 성단 색등급도", "Gaia DR3 카탈로그", "구성원 선별 → 색등급도 → 등시선 맞춤", "선별 엄격도 · 나이 · 거리 · 소광"),
]
tw = W - 2 * M
tbl = sl.shapes.add_table(4, 4, I(M), I(y + 0.10), I(tw), I(2.4)).table
for c, wd in zip(range(4), (3.05, 2.85, 4.30, 3.90)):
    tbl.columns[c].width = I(tw * wd / 14.1)
for r, row in enumerate(rows):
    tbl.rows[r].height = I(0.60 if r else 0.44)
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = I(0.10)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TBLHEAD if r == 0 else BG
        p = cell.text_frame.paragraphs[0]
        p.line_spacing = 1.1
        for run in p.runs:
            run.font.size = Pt(14)
            run.font.name = F
            run.font.bold = (r == 0)
            run.font.color.rgb = WHITE if r == 0 else TBLTEXT
f = tb(sl, M, y + 2.72, tw, 1.1)
put(f, "설치도 로그인도 없다. 브라우저에서 바로 열린다.", 17, BODY, first=True, space_after=8)
put(f, "다만 세부 구현과 검토 범위는 같지 않다. 사용자 검토를 받은 것은 식현상 모듈뿐이다.",
    17, WARN, True)

# ═════ 5. 웹 기능 ① ═══════════════════════════════════════════════════
sl = S()
y = title(sl, "웹 기능 ① 탐구 질문 진입",
          "천체명·좌표가 아닌 질문에서 대상과 자료로 이어진다")
pic(sl, crop("step0_entry.png"), M, y, 7.75, H - y - 0.75, root=HERE)
bullets(sl, M + 8.15, y + 0.2, W - M - 8.15 - M + 0.45, [
    ("첫 화면은 세 모듈을 각각의 대표 탐구 질문과 사용 자료로 보여 준다.", BODY, False),
    ("「행성 크기를 어떻게 알아낼 수 있을까」에서 출발해 대상과 자료로 이어진다.", BODY, False),
    ("자료 확인 단계에서 섹터·케이던스·관측 기간을 먼저 읽게 한다.", BODY, False),
    ("자료를 고르는 일이 이미 탐구의 일부다.", ACC, True),
], size=17, gap=15)

# ═════ 6. 웹 기능 ② ═══════════════════════════════════════════════════
sl = S()
y = title(sl, "웹 기능 ② 분석 조건 설정",
          "구경·배경·비교성을 학습자가 정한다")
pic(sl, crop("step3_conditions.png"), M, y, 7.75, H - y - 0.75, root=HERE, top=True)
bullets(sl, M + 8.15, y + 0.2, W - M - 8.15 - M + 0.45, [
    ("구경 반지름, 배경 고리, 비교성 선택을 학습자가 정한다. 기본값은 2.5픽셀과 4.0~6.0픽셀이다.", BODY, False),
    ("측광·모델 적합은 자동으로 돌지만 조건과 품질 지표는 가리지 않는다.", BODY, False),
    ("조건이 값을 바꾼다는 것 자체가 학습 내용이다.", ACC, True),
], size=17, gap=15)

# ═════ 7. 웹 기능 ③ ═══════════════════════════════════════════════════
sl = S()
y = title(sl, "웹 기능 ③ 기준값 비교와 기록",
          "산출값을 카탈로그 값과 나란히 두고 해석을 남긴다")
pic(sl, "_panel_step5_reference-step6_record.png", M, y + 0.05, W - 2 * M, 3.95)
bullets(sl, M, y + 4.12, W - 2 * M, [
    ("왼쪽 — 위상 접기 광도곡선에 적합 모델과 카탈로그 기대 모델을 겹치고 잔차를 함께 보인다. "
     "기준 반지름비는 아카이브 수록값을 받아 온다.", BODY, False),
    ("오른쪽 — 해석 질문과 서술형 기록 칸을 둔다. 과제는 값을 맞히는 것이 아니라 차이를 설명하는 것이다.",
     ACC, True),
], size=16, gap=9)
cite(sl, "대상 WASP-6 b · 자료 MAST TESScut · 기준값 NASA Exoplanet Archive")

# ═════ 8. 바이브 코딩 ═════════════════════════════════════════════════
sl = S()
y = title(sl, "개발 방식과 검증",
          "AI가 쓴 코드는 실행되어도 값이 틀릴 수 있다")
bullets(sl, M, y, W - 2 * M, [
    ("수업 맥락을 아는 연구자가 자연어로 의도를 적고 생성된 코드를 실행해 확인하는 방식으로 직접 구현하였다.",
     BODY, False),
    ("사용 범위는 프론트엔드·백엔드 코드 초안, 오류 수정, 반복 구현이다. "
     "탐구 구조와 단계 설계, 학습자에게 개방할 분석 조건, 천문 모델과 가정은 연구자가 정했다.", BODY, False),
    ("생성된 코드는 오류 없이 실행되면서도 산출값이 틀릴 수 있다. "
     "화면이 돌아간다고 값을 믿을 수는 없다.", WARN, True),
    ("", BODY, False),
    ("그래서 두 가지를 따로 확인했다 — 산출값은 문헌값과 대조하고, 화면의 안내 문장은 사용자 검토로 점검했다.",
     ACC, True),
    ("개발 방식 간 비교는 하지 않았다. 시간과 비용을 기록하지 않아 빨랐다고 말할 수 없다.", GREY, False),
], size=18, gap=13)
cite(sl, "Michels et al. (2026) · Song et al. (2026) · Uddin (2026)")

# ═════ 9. 결과 ① 검증 ═════════════════════════════════════════════════
sl = S()
y = title(sl, "WASP-121 b 처리 조건별 결과",
          "-12.8%가 -2.8%까지 좁혀졌다")
f = tb(sl, M, y - 0.08, W - 2 * M, 0.34)
put(f, "같은 설정을 반복 실행하면 WASP-6 b 반지름비가 0.14534로 재현된다. "
       "재현성과 정확도는 다른 문제다.", 13, GREY, first=True)
y += 0.30
pic(sl, "fig_table7.png", M, y - 0.02, W - 2 * M, H - y - 0.62, root=HERE)
cite(sl, "Daylan et al. (2021)과 같은 자료(TESS 섹터 7 · 2분 케이던스)를 별도 스크립트로 분석. "
         "플랫폼의 전체 실행 경로와는 다르다.")

# ═════ 10. 결과 ② 한계 ════════════════════════════════════════════════
sl = S()
y = title(sl, "점검의 한계",
          "남은 -2.8%와 비교성의 효과는 나누지 못했다")
bullets(sl, M, y + 0.1, W - 2 * M, [
    ("다섯 조건은 서로 다른 처리를 여러 개 함께 포함한다. 표의 순서대로 차이가 누적되지 않고, "
     "한 요인의 효과로도 읽을 수 없다.", BODY, False),
    ("남은 -2.8%에는 기준선 처리, 모델 설정, 적합 방법 등 통제하지 않은 차이가 함께 들어 있다.", BODY, False),
    ("표준 광도곡선 파일의 화소 범위에 비교성이 없어, 비교성의 수와 종류에 따른 변화는 확인하지 못했다.",
     BODY, False),
    ("", BODY, False),
    ("교육적으로는 이 민감도가 오히려 자산이다. 조건을 드러내야 차이를 설명하는 활동이 성립한다.",
     ACC, True),
], size=18, gap=14)

# ═════ 11. 결과 ③ 사용자 검토 ═════════════════════════════════════════
# 그림 둘은 다른 세션이 논문용으로 만든 정본이다(docs/make_survey_figs.py).
# 발표와 논문이 같은 그림을 쓰도록 그것을 그대로 가져온다. 다크판은 색만 바꾼 사본이다.
sl = S()
y = title(sl, "사용자 검토 결과",
          "실행 부담은 낮고 기준값 해석이 최하위였다")
pic(sl, "fig_survey_likert.png", M, y - 0.02, W - 2 * M, H - y - 0.62, root=HERE)
cite(sl, "1차 현직교사 중심 13명 (2026-07-24) · 2차 예비교사 13명 (2026-09-06~07). "
         "5점 척도. * 는 역채점한 부정 진술.")

# ═════ 12. 현장 적용 ══════════════════════════════════════════════════
sl = S()
y = title(sl, "보완 요구",
          "용어·그래프 해석과 수업용 자료에 몰렸다")
pic(sl, "fig_survey_needs.png", M, y - 0.05, W - 2 * M, H - y - 0.58, root=HERE)
cite(sl, "(a) 2차 예비교사 13명 · (b) 두 조사의 보완 요구(복수선택).")

# ═════ 13. 결론 ═══════════════════════════════════════════════════════
sl = S()
y = title(sl, "결론",
          "조건을 드러내야 차이를 설명하는 활동이 성립한다")
bullets(sl, M, y + 0.15, W - 2 * M, [
    ("공개 천문 아카이브의 세 자료를 같은 일곱 단계 탐구 흐름에 올려, 코딩 없이 웹에서 "
     "분석하도록 구현하였다.", BODY, False),
    ("생성형 AI 코딩 도구로 만들었기 때문에 산출값을 따로 검증했다. 반복 실행에서 재현되었고, "
     "문헌값과의 차이는 처리 조건으로 -12.8%에서 -2.8%까지 설명되었다.", BODY, False),
    ("교사 26명 검토에서 실행 부담 완화는 확인되었고, 기준값 비교 화면의 해석 지원이 "
     "다음 과제로 남았다.", BODY, False),
    ("", BODY, False),
    ("실행을 돕는 일과 해석을 돕는 일은 다른 과제다.", ACC, True),
], size=18, gap=14)
rule(sl, H - 1.38)
f = tb(sl, M, H - 1.20, W - 2 * M, 0.9)
put(f, "https://easwa-webapp.onrender.com", 18, ACC, True, first=True, space_after=5)
put(f, "박민준 · 한국교원대학교 지구과학교육과 · pmj3265@gmail.com", 13, GREY)

# ═════ 14. 참고문헌 ═══════════════════════════════════════════════════
sl = S()
y = title(sl, "참고문헌")
refs = [
    "Claret A (2017) Limb and gravity-darkening coefficients for the TESS satellite. A&A 600: A30.",
    "Daylan T, Gunther M N, Mikal-Evans T, et al. (2021) TESS observations of the WASP-121 b phase curve. AJ 161: 131.",
    "Delrez L, Santerne A, Almenara J-M, et al. (2016) High-precision multi-wavelength eclipse photometry of WASP-121 b. MNRAS 458: 4025-4043.",
    "Gaia Collaboration, Vallenari A, Brown A G A, et al. (2023) Gaia Data Release 3. A&A 674: A1.",
    "Kreidberg L (2015) batman: BAsic Transit Model cAlculatioN in Python. PASP 127: 1161-1165.",
    "Mandel K and Agol E (2002) Analytic light curves for planetary transit searches. ApJ 580: L171-L175.",
    "Michels D L, Abu Ghazaleh M, Lazzari F, Kassem N and Klein J (2026) Vibe coding: Practice, performance, productivity, and risk. arXiv:2608.20446.",
    "Paczynski B (1986) Gravitational microlensing by the galactic halo. ApJ 304: 1-5.",
    "Song Y, Choi S, Kim J, Kim Y, Weisberg L and Moon J (2026) A guiding framework for K-12 teachers in creating AI-powered learning technologies through vibe coding. arXiv:2607.05406.",
    "Uddin S M J (2026) Is vibe coding the future? An empirical assessment of LLM generated codes for construction safety. arXiv:2604.12311.",
    "Wong N, Elsayed R, Perez L R, Nilsen K, Daehler K R and Darche S (2026) Data-rich science instruction. Education Sciences 16: 171.",
    "교육부 (2022) 과학과 교육과정. 교육부 고시 제2022-33호 [별책 9].",
]
f = tb(sl, M, y, W - 2 * M, H - y - 0.6)
for i, r in enumerate(refs):
    put(f, r, 13, BODY, space_after=6, first=(i == 0), line=1.1)

os.makedirs(DEST, exist_ok=True)
prs.save(OUT)
print("저장 — %s" % OUT)
print("  %d장 · %.1f MB" % (len(prs.slides._sldIdLst), os.path.getsize(OUT) / 1e6))
